#!/usr/bin/env python3
"""Supplementary-data rescue pilot (NCI ODS Impact Prize, Track 1). Written 2026-09-22.

Question: of NCI-funded open-access papers on one cancer type (2016-2022) that have supplementary
files but cite no data repository, how many of those files are real, reusable datasets?

Steps
  sample    pick N random papers via Europe PMC (seeded), save sample.json
  fetch     download each paper's supplementary files (Europe PMC supplementaryFiles endpoint)
  classify  open every file, measure tables, classify, write files.csv + papers.csv
  report    print summary tables

Everything uses public APIs; no credentials.
"""
import argparse, csv, io, json, os, random, re, statistics, sys, time, zipfile, urllib.request, urllib.parse, collections

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(HERE, "data")
EPMC = "https://www.ebi.ac.uk/europepmc/webservices/rest"
UA = {"User-Agent": "ods-rescue-pilot/1.0 (research; contact via github)"}
ACC = "(ACCESSION_TYPE:gen OR ACCESSION_TYPE:sra OR ACCESSION_TYPE:dbgap OR ACCESSION_TYPE:ega OR ACCESSION_TYPE:arrayexpress OR ACCESSION_TYPE:pride OR ACCESSION_TYPE:bioproject OR ACCESSION_TYPE:metabolights OR ACCESSION_TYPE:pdb)"
MAX_ZIP_MB = 120


def get(url, timeout=90, binary=False, cap_mb=None):
    req = urllib.request.Request(url, headers=UA)
    for attempt in range(4):
        try:
            with urllib.request.urlopen(req, timeout=timeout) as r:
                if cap_mb:
                    buf = io.BytesIO(); n = 0
                    while True:
                        chunk = r.read(1 << 20)
                        if not chunk: break
                        n += len(chunk)
                        if n > cap_mb << 20: return None, "too_large"
                        buf.write(chunk)
                    return buf.getvalue(), r.headers.get("Content-Type", "")
                b = r.read()
                return (b if binary else b.decode("utf-8", "replace")), r.headers.get("Content-Type", "")
        except Exception as e:
            err = str(e)[:80]; time.sleep(2 + 3 * attempt)
    return None, err


def cmd_sample(a):
    os.makedirs(DATA, exist_ok=True)
    q = (f'GRANT_AGENCY:"NCI NIH HHS" AND PUB_YEAR:[2016 TO 2022] AND OPEN_ACCESS:y AND HAS_SUPPL:Y '
         f'AND NOT {ACC} AND (TITLE:"{a.cancer}" OR ABSTRACT:"{a.cancer}")')
    ids, cursor = [], "*"
    while True:
        u = f"{EPMC}/search?query={urllib.parse.quote(q)}&format=json&pageSize=1000&resultType=lite&cursorMark={urllib.parse.quote(cursor)}"
        t, _ = get(u); d = json.loads(t)
        res = d.get("resultList", {}).get("result", [])
        for r in res:
            if r.get("pmcid"):
                ids.append({k: r.get(k) for k in ("pmcid", "pmid", "doi", "title", "journalTitle", "pubYear", "license", "authorString")})
        nxt = d.get("nextCursorMark")
        if not res or not nxt or nxt == cursor: break
        cursor = nxt
    print(f"query hit count {d.get('hitCount')} ; collected {len(ids)} with a PMC id")
    random.seed(a.seed); random.shuffle(ids)
    sample = ids[: a.n]
    json.dump({"cancer": a.cancer, "query": q, "population": len(ids), "seed": a.seed, "papers": sample},
              open(os.path.join(DATA, "sample.json"), "w"), indent=1)
    print(f"sampled {len(sample)} -> data/sample.json")


def cmd_fetch(a):
    import concurrent.futures as cf, threading
    s = json.load(open(os.path.join(DATA, "sample.json")))
    zd = os.path.join(DATA, "zips"); os.makedirs(zd, exist_ok=True)
    lp = os.path.join(DATA, "fetch_log.json")
    log = json.load(open(lp)) if os.path.exists(lp) else {}
    todo = [p["pmcid"] for p in s["papers"] if log.get(p["pmcid"], {}).get("status") not in ("ok", "empty", "too_large")]
    lock = threading.Lock(); done = [0]
    def one(pmc):
        b, ct = get(f"{EPMC}/{pmc}/supplementaryFiles?includeInlineImage=false", timeout=300, binary=True, cap_mb=MAX_ZIP_MB)
        if b is None: st = {"status": "too_large" if ct == "too_large" else "error", "note": ct}
        elif "zip" in (ct or "") and len(b) > 100:
            open(os.path.join(zd, pmc + ".zip"), "wb").write(b); st = {"status": "ok", "bytes": len(b)}
        else: st = {"status": "empty", "note": (ct or "")[:40]}
        with lock:
            log[pmc] = st; done[0] += 1
            json.dump(log, open(lp, "w"), indent=0)
            print(f"[{done[0]}/{len(todo)}] {pmc} {st['status']} {st.get('bytes','')}", flush=True)
    with cf.ThreadPoolExecutor(a.workers) as ex: list(ex.map(one, todo))
    print(dict(collections.Counter(v["status"] for v in log.values())))


# ---------- classification ----------
TAB_EXT = {".xlsx", ".xlsm", ".xls", ".csv", ".tsv", ".txt", ".tab"}
DOC_EXT = {".docx", ".doc", ".pdf", ".rtf"}
IMG_EXT = {".tif", ".tiff", ".jpg", ".jpeg", ".png", ".gif", ".eps", ".pptx", ".ppt", ".ai", ".svg", ".bmp", ".emf", ".psd"}
MEDIA_EXT = {".mp4", ".avi", ".mov", ".wmv", ".mpg", ".mpeg", ".mp3", ".wav", ".gif"}
CODE_EXT = {".r", ".py", ".m", ".rmd", ".ipynb", ".sh", ".pl", ".java", ".cpp", ".jl", ".sas", ".do"}
GENE = re.compile(r"^[A-Z][A-Z0-9orf\-]{1,9}$")
TCGA = re.compile(r"^TCGA-[A-Z0-9]{2}-[A-Z0-9]{4}")
CLIN = re.compile(r"\b(age|sex|gender|race|stage|grade|survival|os|pfs|dfs|status|vital|death|recurr|follow.?up|treatment|therapy|dose|response|karnofsky|ecog|smok|bmi|idh|mgmt|tumou?r size|histolog)\b", re.I)
DERIV = re.compile(r"\b(GO:\d+|KEGG|reactome|enrich|pathway|-?log10?\s*\(?\s*p|padj|p\.?adj|fdr|q.?value|fold.?change|log2 ?fc|DEG|differentially|hallmark|gsea|NES)\b", re.I)
SEQ = re.compile(r"\b(fpkm|tpm|rpkm|log2 ?fc|fold.?change|p.?adj|padj|fdr|q.?value|count|reads|logcpm|cpm|base ?mean|zscore|z-score|ensembl|entrez|gene.?symbol|gene.?id|peptide|protein|uniprot|m/z|intensity|beta.?value|cg\d{5,})\b", re.I)


def table_stats(rows):
    rows = [r for r in rows if any(str(c).strip() for c in r if c is not None)]
    if not rows: return None
    w = max(len(r) for r in rows); n = len(rows)
    cells = [str(c).strip() for r in rows for c in r if c is not None and str(c).strip() != ""]
    if not cells: return None
    def isnum(x):
        try: float(x.replace(",", "")); return True
        except: return False
    numfrac = sum(isnum(c) for c in cells) / len(cells)
    header = " ".join(str(c) for c in rows[0] if c is not None)
    first_col = [str(r[0]).strip() for r in rows[1:] if r and r[0] is not None]
    genefrac = sum(bool(GENE.match(x)) for x in first_col) / max(len(first_col), 1)
    tcga = sum(bool(TCGA.match(x)) for x in first_col)
    return {"rows": n, "cols": w, "numfrac": round(numfrac, 2), "genefrac": round(genefrac, 2), "tcga_ids": tcga,
            "clin_terms": len(set(m.lower() for m in CLIN.findall(header))), "derived": bool(DERIV.search(header) or DERIV.search(" ".join(first_col[:5]))), "seq_terms": len(set(m.lower() for m in SEQ.findall(header))),
            "header": header[:200]}


def read_tables(name, b):
    ext = os.path.splitext(name.lower())[1]; out = []
    try:
        if ext in (".xlsx", ".xlsm"):
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(b), read_only=True, data_only=True)
            for ws in wb.worksheets:
                rows = []
                for i, r in enumerate(ws.iter_rows(values_only=True)):
                    rows.append(r)
                    if i > 200000: break
                st = table_stats(rows)
                if st: st["sheet"] = ws.title; out.append(st)
        elif ext == ".xls":
            import xlrd
            wb = xlrd.open_workbook(file_contents=b)
            for ws in wb.sheets():
                rows = [ws.row_values(i) for i in range(min(ws.nrows, 200000))]
                st = table_stats(rows)
                if st: st["sheet"] = ws.name; out.append(st)
        elif ext in (".csv", ".tsv", ".txt", ".tab"):
            t = b[: 50 << 20].decode("utf-8", "replace")
            delim = "\t" if t.count("\t") > t.count(",") else ","
            rows = list(csv.reader(io.StringIO(t), delimiter=delim))
            st = table_stats(rows)
            if st and st["cols"] >= 2: st["sheet"] = "-"; out.append(st)
        elif ext == ".docx":
            import docx
            d = docx.Document(io.BytesIO(b))
            for k, t in enumerate(d.tables):
                rows = [[c.text for c in r.cells] for r in t.rows]
                st = table_stats(rows)
                if st and st["cols"] >= 3 and st["numfrac"] >= 0.3: st["sheet"] = f"table{k+1}"; out.append(st)
        elif ext == ".pdf":
            import pdfplumber, logging
            logging.getLogger("pdfminer").setLevel(logging.ERROR)
            with pdfplumber.open(io.BytesIO(b)) as pdf:
                pages = pdf.pages[:60]
                first = (pages[0].extract_text() or "")[:400] if pages else ""
                if re.search(r"reviewer|peer review|response to review|rebuttal|graphical abstract|disclaims all liability", first, re.I):
                    return [{"error": "not_data_pdf"}]
                # keep only real grids (>=3 columns, >=5 rows); merge across pages
                allrows = []
                for pg in pages:
                    for tb in pg.extract_tables() or []:
                        rows = [[(c or "") for c in r] for r in tb if r]
                        if rows and len(rows) >= 5 and max(len(r) for r in rows) >= 3: allrows += rows
                if allrows:
                    st = table_stats(allrows)
                    if st and st["numfrac"] >= 0.3: st["sheet"] = f"pdf_tables({len(pages)}p)"; out.append(st)
    except Exception as e:
        return [{"error": str(e)[:60]}]
    return out


def classify_file(name, size, tables):
    ext = os.path.splitext(name.lower())[1]
    if ext in CODE_EXT: return "code", ""
    if ext in MEDIA_EXT and ext != ".gif": return "media", ""
    if ext in IMG_EXT: return "figure_image", ""
    good = [t for t in tables if "rows" in t]
    if ext in TAB_EXT or (ext in (".docx", ".pdf") and good):
        if not good: return ("unreadable_table" if any("error" in t for t in tables) else "empty_or_text"), ""
        big = max(good, key=lambda t: t["rows"] * t["cols"])
        kind = ("clinical" if big["clin_terms"] >= 3 or big["tcga_ids"] > 5 else
                "omics" if big["genefrac"] > 0.5 or big["seq_terms"] >= 2 else
                "numeric" if big["numfrac"] > 0.5 else "mixed")
        total_rows = sum(t["rows"] for t in good) if ext in TAB_EXT else big["rows"]
        if big["rows"] >= 1000 or total_rows >= 2000: return "dataset_large", kind
        if big["rows"] >= 50 and big["cols"] >= 3: return "dataset_medium", kind
        if big["rows"] >= 10 and big["cols"] >= 2: return "small_table", kind
        return "tiny_table", kind
    if ext in DOC_EXT: return "document", ""
    return "other", ""


def cmd_classify(a):
    s = json.load(open(os.path.join(DATA, "sample.json"))); log = json.load(open(os.path.join(DATA, "fetch_log.json")))
    meta = {p["pmcid"]: p for p in s["papers"]}
    frows = []
    for pmc, st in log.items():
        if st["status"] != "ok": continue
        try: z = zipfile.ZipFile(os.path.join(DATA, "zips", pmc + ".zip"))
        except Exception: continue
        for zi in z.infolist():
            if zi.is_dir(): continue
            name = zi.filename; ext = os.path.splitext(name.lower())[1]
            b = z.read(zi) if (ext in TAB_EXT or ext in (".docx", ".pdf")) and zi.file_size < (40 << 20) else b""
            tables = read_tables(name, b) if b else []
            cls, kind = classify_file(name, zi.file_size, tables)
            good = [t for t in tables if "rows" in t]
            big = max(good, key=lambda t: t["rows"] * t["cols"]) if good else {}
            frows.append({"pmcid": pmc, "journal": meta[pmc]["journalTitle"], "year": meta[pmc]["pubYear"], "license": meta[pmc].get("license", ""),
                          "file": os.path.basename(name), "ext": ext, "bytes": zi.file_size, "class": cls, "kind": kind,
                          "native": ext in TAB_EXT, "n_tables": len(good), "max_rows": big.get("rows", ""), "max_cols": big.get("cols", ""),
                          "numfrac": big.get("numfrac", ""), "genefrac": big.get("genefrac", ""), "tcga_ids": big.get("tcga_ids", ""), "derived": big.get("derived", ""),
                          "header": big.get("header", "")})
    with open(os.path.join(DATA, "files.csv"), "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(frows[0].keys())); w.writeheader(); w.writerows(frows)
    # per paper
    prow = []
    for pmc, p in meta.items():
        fs = [r for r in frows if r["pmcid"] == pmc]
        cls = collections.Counter(r["class"] for r in fs)
        best = ("dataset_large" if cls["dataset_large"] else "dataset_medium" if cls["dataset_medium"] else
                "small_table" if cls["small_table"] else "tables_only" if (cls["tiny_table"] or cls["unreadable_table"]) else
                "no_supp_retrieved" if log.get(pmc, {}).get("status") != "ok" else "no_tabular_data")
        prow.append({"pmcid": pmc, "journal": p["journalTitle"], "year": p["pubYear"], "license": p.get("license", ""), "fetch": log.get(pmc, {}).get("status"),
                     "n_files": len(fs), "n_dataset_files": cls["dataset_large"] + cls["dataset_medium"], "n_small_tables": cls["small_table"],
                     "n_figures": cls["figure_image"], "n_docs": cls["document"], "verdict": best,
                     "kinds": ";".join(sorted(set(r["kind"] for r in fs if r["kind"] and r["class"].startswith("dataset")))), "title": p["title"]})
    with open(os.path.join(DATA, "papers.csv"), "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(prow[0].keys())); w.writeheader(); w.writerows(prow)
    print(f"classified {len(frows)} files across {len(prow)} papers -> data/files.csv, data/papers.csv")


def cmd_report(a):
    s = json.load(open(os.path.join(DATA, "sample.json")))
    P = list(csv.DictReader(open(os.path.join(DATA, "papers.csv")))); F = list(csv.DictReader(open(os.path.join(DATA, "files.csv"))))
    n = len(P); ok = [p for p in P if p["fetch"] == "ok"]
    print(f"\nCANCER: {s['cancer']} | population {s['population']} papers | sample {n} | supplements retrieved for {len(ok)}")
    print("\nPER PAPER (what is the best thing in its attachments)")
    for k, v in collections.Counter(p["verdict"] for p in P).most_common(): print(f"  {k:20} {v:4}  {v/n:5.1%}")
    ds = [p for p in P if p["verdict"].startswith("dataset")]
    print(f"\n  => papers with at least one real dataset file (>=50 rows x 3 cols, or >=1000 rows): {len(ds)} of {n} = {len(ds)/n:.1%}")
    print(f"  => of papers whose supplements were retrieved: {len([p for p in ok if p['verdict'].startswith('dataset')])/max(len(ok),1):.1%}")
    print("\nPER FILE")
    for k, v in collections.Counter(f["class"] for f in F).most_common(): print(f"  {k:20} {v:5}  {v/len(F):5.1%}")
    dsf = [f for f in F if f["class"].startswith("dataset")]
    print(f"\nDATASET FILES: {len(dsf)}; kinds:", dict(collections.Counter(f['kind'] for f in dsf)))
    rows = [int(f["max_rows"]) for f in dsf if f["max_rows"]]
    if rows: print(f"  largest-table rows: median {statistics.median(rows):.0f}, p90 {sorted(rows)[int(.9*len(rows))]}, max {max(rows)}")
    print("  by extension:", dict(collections.Counter(f["ext"] for f in dsf)))
    nat = {f["pmcid"] for f in dsf if f["native"] == "True"}; docd = {f["pmcid"] for f in dsf if f["native"] != "True"} - nat
    print(f"  papers with a NATIVE spreadsheet/csv dataset: {len(nat)} ({len(nat)/n:.1%} of sample) | dataset only as a table inside a PDF/DOCX: {len(docd)} ({len(docd)/n:.1%})")
    der = {f["pmcid"] for f in dsf if f["derived"] == "True"}; prim = {f["pmcid"] for f in dsf if f["derived"] != "True"}
    print(f"  papers with a PRIMARY-looking dataset (measurements, cohort/sample tables): {len(prim)} ({len(prim)/n:.1%}) | only DERIVED result tables (enrichment, DEG lists, stats): {len(der - prim)} ({len(der-prim)/n:.1%})")
    lic = collections.Counter((p["license"] or "untagged") for p in ds)
    print("  licenses of dataset papers:", dict(lic), f"| freely re-depositable (cc by / cc0): {sum(v for k,v in lic.items() if k in ('cc by','cc0'))} of {len(ds)}")
    print("\nLICENSES of sampled papers:", dict(collections.Counter(p["license"] or "unknown" for p in P)))
    print("\nTOP JOURNALS among dataset papers:", collections.Counter(p["journal"] for p in ds).most_common(8))
    print(f"\nEXTRAPOLATION (rough): {len(ds)/n:.1%} of {s['population']} {s['cancer']} papers ~= {round(len(ds)/n*s['population'])} papers with rescuable datasets in this cancer alone;")
    print(f"  applied to the ~32,000 NCI-funded papers with supplements and no accession -> ~{round(len(ds)/n*32000):,} papers (upper-bound heuristic, needs validation)")


if __name__ == "__main__":
    ap = argparse.ArgumentParser(); sub = ap.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("sample"); s.add_argument("--cancer", default="glioblastoma"); s.add_argument("--n", type=int, default=300); s.add_argument("--seed", type=int, default=20260922); s.set_defaults(fn=cmd_sample)
    f = sub.add_parser("fetch"); f.add_argument("--workers", type=int, default=6); f.set_defaults(fn=cmd_fetch); sub.add_parser("classify").set_defaults(fn=cmd_classify); sub.add_parser("report").set_defaults(fn=cmd_report)
    a = ap.parse_args(); a.fn(a)
